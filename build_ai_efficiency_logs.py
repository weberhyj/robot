from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "ai_log_render"
SCREENSHOT_DIR = OUT_DIR / "screenshots"
OUT_FILE = OUT_DIR / "AI增效日志_ROBOFLOW前端开发_汇总版.docx"

COLORS = {
    "ink": "111827",
    "muted": "667085",
    "blue": "2563EB",
    "green": "0F9F6E",
    "red": "E5484D",
    "line": "D8DEE8",
    "soft": "F7F9FC",
    "header": "EAF1FF",
}


def set_run_font(run, size: float = 10.5, bold: bool = False, color: str = "111827") -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc: Document, text: str, size: float = 10.5, color: str = "111827", bold: bool = False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.2
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    return paragraph


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13, bold=True, color=COLORS["blue"])
    return paragraph


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Cm(0.7)
        paragraph.paragraph_format.first_line_indent = Cm(-0.25)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        set_run_font(run, size=10.2, color=COLORS["ink"])


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D8DEE8", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 90, start: int = 130, bottom: int = 90, end: int = 130) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_cm: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width / 2.54 * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def style_table(table, header_rows: int = 1) -> None:
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_index < header_rows:
                set_cell_shading(cell, COLORS["header"])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.12
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=9.4 if row_index >= header_rows else 9.8,
                        bold=row_index < header_rows,
                        color=COLORS["ink"],
                    )


def add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_widths(table, [4.0, 12.6])
    for index, (key, value) in enumerate(rows):
        table.cell(index, 0).text = key
        table.cell(index, 1).text = value
    style_table(table, header_rows=0)
    for row in table.rows:
        set_cell_shading(row.cells[0], COLORS["soft"])
        for run in row.cells[0].paragraphs[0].runs:
            set_run_font(run, size=9.5, bold=True, color=COLORS["muted"])


def add_matrix(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_widths(table, widths_cm)
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value
    style_table(table, header_rows=1)


def add_screenshot(doc: Document, image_name: str, caption: str) -> None:
    image_path = SCREENSHOT_DIR / image_name
    if not image_path.exists():
        add_text(doc, f"截图缺失：{image_path}", color=COLORS["red"])
        return

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(16.2))

    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, size=9, color=COLORS["muted"])


def add_log(doc: Document, log: dict[str, object]) -> None:
    add_heading(doc, str(log["title"]), level=1)
    add_kv_table(
        doc,
        [
            ("项目", "ROBOFLOW 智能视觉上料与机器人装配工作站"),
            ("姓名", "[填写姓名]"),
            ("所属工种", "前端开发工程师 / UI 工程师"),
            ("对应模块", str(log["module"])),
            ("开发任务名称", str(log["task_name"])),
            ("任务开始时间", "[YYYY-MM-DD HH:MM，按实际填写]"),
            ("任务结束时间", "[YYYY-MM-DD HH:MM，按实际填写]"),
            ("AI工具清单", "Codex / ChatGPT，Figma 参考链接，本地 Vue 前端工程工具链"),
        ],
    )

    add_heading(doc, "任务描述", level=2)
    add_text(doc, str(log["description"]))
    add_bullets(doc, log["features"])  # type: ignore[arg-type]

    add_heading(doc, "传统做法预估", level=2)
    add_text(doc, f"如果不用 AI，我估计需要：{log['estimated']} 小时。")
    add_matrix(
        doc,
        ["传统环节", "耗时预估", "说明"],
        log["traditional"],  # type: ignore[arg-type]
        [4.5, 2.5, 9.6],
    )

    add_heading(doc, "关键 Prompt 与迭代", level=2)
    add_text(doc, "以下为本次对话中可复用的关键需求表达，已整理成适合归档的版本。", color=COLORS["muted"])
    add_matrix(
        doc,
        ["轮次", "Prompt / 需求表达", "AI 输出或迭代结果"],
        log["prompts"],  # type: ignore[arg-type]
        [1.4, 7.1, 8.1],
    )

    add_heading(doc, "效率与质量总结", level=2)
    estimated = float(log["estimated"])
    actual = float(log["actual"])
    efficiency = round((estimated - actual) / estimated * 100, 1)
    add_kv_table(
        doc,
        [
            ("实际最终耗时", f"{actual} 小时，可按真实投入时间修正"),
            ("效率提升", f"[({estimated} - {actual}) / {estimated} × 100%] = {efficiency}%"),
            ("成果质量自评", str(log["quality"])),
            ("提效最明显环节", str(log["biggest_gain"])),
        ],
    )

    add_heading(doc, "可复用的模式与建议", level=2)
    add_text(doc, "万能 Prompt 模式：", bold=True)
    add_text(doc, str(log["prompt_pattern"]), color=COLORS["blue"])
    add_text(doc, "给类似任务工程师的建议：", bold=True)
    add_bullets(doc, log["suggestions"])  # type: ignore[arg-type]

    screenshots = log.get("screenshots", [])
    if screenshots:
        add_heading(doc, "截图附件", level=2)
        for image_name, caption in screenshots:  # type: ignore[assignment]
            add_screenshot(doc, image_name, caption)

    add_heading(doc, "附件清单", level=2)
    add_matrix(
        doc,
        ["附件类型", "建议归档内容"],
        [
            ["代码与配置", str(log["code_assets"])],
            ["页面截图", str(log["screen_assets"])],
            ["验证记录", str(log["verification"])],
        ],
        [4.2, 12.4],
    )


LOGS: list[dict[str, object]] = [
    {
        "title": "日志 1：项目规则、目录架构与开发规范沉淀",
        "module": "项目规范 / 前端架构 / 协作规则",
        "task_name": "沉淀 AGENTS.md、项目目录架构、路由规范、组件规范和前端开发边界",
        "description": "围绕 ROBOFLOW 上位机机器人项目，先明确当前阶段只开发前端，再把长期有效的工程规则、技术栈边界、目录分层和组件协作方式整理成文档，避免页面迭代时规则分散。",
        "features": [
            "将项目边界明确为 frontend-only，暂不接入真实机器人、PLC、相机、串口、Electron 主进程、安装包和真实持久化。",
            "建立 Vue 3、Vite、TypeScript、Pinia、Vue Router、Vue I18n、Element Plus、ECharts、SCSS、pnpm、Vitest、@antfu/eslint-config 的默认技术标准。",
            "沉淀 AGENTS.md 纯规则版，排除页面规格、接口、类型定义和开发计划等易变化内容。",
            "整理项目目录架构、路由规范、组件规范、mock 数据契约和数据流文档，便于后续多人维护。",
        ],
        "estimated": 6.0,
        "actual": 1.2,
        "traditional": [
            ["需求边界梳理", "1.0 h", "人工逐条确认哪些属于前端阶段，哪些需要等 EXE 或设备集成阶段再做。"],
            ["技术栈与规范收敛", "1.5 h", "结合项目现状、Antfu 习惯和团队偏好，人工整理长期规则。"],
            ["目录与路由规范编写", "2.0 h", "设计 pages、widgets、entities、stores、services、mocks、shared 的职责边界。"],
            ["规则去重与文档清理", "1.5 h", "删除计划、接口、验收清单等不适合放进 AGENTS.md 的内容。"],
        ],
        "prompts": [
            ["1", "先输出该项目的 AGENTS.md，并追加 pnpm、Element Plus、SCSS、ECharts 等规则。", "生成初版项目规则，明确前端命令、组件库、图表库和样式方案。"],
            ["2", "该文件不要中英混合，用全部英文吧；除了规则，其他东西不要放进这里。", "将 AGENTS.md 收敛为英文纯规则文档，剔除计划和页面规格。"],
            ["3", "路由规格如何规范？项目目录架构呢？", "补充 route-spec、project-structure 等工程规范文档。"],
            ["4", "补充 component-conventions.md。", "沉淀组件命名、职责、复用、样式和数据输入规范。"],
        ],
        "quality": "完整性 5 分；可维护性 5 分；团队复用性 5 分。",
        "biggest_gain": "AI 对规则去重、结构化归类和英文规范表达提效最明显，减少了从零写项目规范的时间。",
        "prompt_pattern": "请基于[项目阶段]和[已有技术栈]，只输出长期有效的[规则类型]，不要包含[页面规格/接口/计划/验收清单]，并按[团队偏好]统一语言和命令规范。",
        "suggestions": [
            "先让 AI 划分“规则”和“计划”，再决定哪些内容写入 AGENTS.md，可以避免规范文档膨胀。",
            "对团队规则要明确优先级，例如项目规则优先于通用 skill 或外部最佳实践。",
            "文档类任务适合用多轮小范围修正，先搭骨架，再逐条合并、删除和改写。",
        ],
        "code_assets": "AGENTS.md，docs/project-structure.md，docs/route-spec.md，docs/component-conventions.md 等。",
        "screen_assets": "无页面截图要求，建议归档文档变更截图或提交记录。",
        "verification": "人工复读 AGENTS.md，确认英文-only、只包含规则、不包含接口和开发计划。",
    },
    {
        "title": "日志 2：仪表盘 Dashboard 页面设计与多轮 UI/UX 优化",
        "module": "仪表盘 / 工作站状态 / 快捷控制 / 历史任务与报警",
        "task_name": "设计并优化 ROBOFLOW 仪表盘，形成紧凑的工业控制台首页",
        "description": "以机器人分拣工作站为主场景，围绕工作站状态、分拣指标、设备状态、快捷控制、历史任务和报警信息进行多轮页面迭代，最终形成浅色、紧凑、可用于工控监控的 dashboard。",
        "features": [
            "status-panel 只保留机器人状态，并将 OK 数、NG 数、识别准确率、平均耗时等指标放到右侧统计区。",
            "快捷控制模块放到底部，包含操作模式切换、启动、停止、复位和急停。",
            "设备状态模块展示机械臂、视觉相机、输送带、夹爪和 A/B/NG 料盒，并包含图片、设备编号、状态和查看详情入口。",
            "历史任务列表和报警信息调整为左右分栏，删除副标题，收窄卡片标题间距，避免纵向滚动条。",
            "统一 dashboard 小卡片 shadow 为 0px 18px 34px -12px rgba(31,43,66, 5.5%)，并去掉小卡片 border。",
        ],
        "estimated": 8.0,
        "actual": 2.0,
        "traditional": [
            ["信息架构设计", "1.5 h", "确定状态、统计、设备、任务和报警的优先级。"],
            ["页面布局实现", "2.5 h", "手工调整顶部状态栏、主卡片、设备横向滚动和底部控制区域。"],
            ["视觉细节优化", "2.5 h", "反复调整阴影、间距、标题、卡片边框和模块高度。"],
            ["截图回归检查", "1.5 h", "检查 1366×768 和 1920×1080 下是否出现溢出或滚动条。"],
        ],
        "prompts": [
            ["1", "status-panel 只保留机器人状态，然后把分拣总数里 OK 数、NG 数、识别准确率、平均耗时放在右边。", "重组顶部核心状态区，弱化冗余信息，突出生产指标。"],
            ["2", "快捷控制按钮放在底部，并添加操作模式：自动或者手动。", "将控制区移动到底部，加入自动/手动切换，符合工控操作路径。"],
            ["3", "效率与分布模块去掉，然后添加各设备的状态模块；每个设备都有图片、设备编号和查看详情按钮。", "替换为更贴合设备监控的横向设备卡模块。"],
            ["4", "所有模块都不要副标题；历史任务列表和报警信息左右分栏。", "压缩界面高度，减少信息噪声，提升首页扫描效率。"],
            ["5", "shadow 改成指定值，而且只针对 *__main-card 里面的小卡片。", "统一局部小卡片阴影，保留最大容器的克制外观。"],
        ],
        "quality": "功能完整性 5 分；视觉一致性 4.5 分；工业控制语义 4.5 分。",
        "biggest_gain": "AI 在多轮局部 UI 微调中提效最明显，能快速定位卡片、间距、阴影、模块顺序和溢出问题。",
        "prompt_pattern": "请在[页面名称]中，把[模块/字段]调整到[目标位置]，保持[现有设计风格]，并满足[不滚动/不溢出/不嵌套卡片]等约束。",
        "suggestions": [
            "对 UI 任务要直接指出具体类名或模块名，例如 status-panel、top-status-bar、device-status-card。",
            "尺寸问题要给出明确结果约束，例如不要超出 panel/main-card，不要出现纵向滚动条。",
            "视觉迭代时先统一标题、间距、阴影和 border，再处理更细的字段对齐。",
        ],
        "screenshots": [("01-dashboard.png", "图 1：ROBOFLOW 仪表盘页面截图")],
        "code_assets": "src/widgets/dashboard-overview/*，src/widgets/top-status-bar/*，src/layouts/WorkbenchLayout.vue，src/shared/styles/main.scss。",
        "screen_assets": "ai_log_render/screenshots/01-dashboard.png。",
        "verification": "已通过本地页面截图检查仪表盘主体内容、设备状态、历史任务和报警区展示。",
    },
    {
        "title": "日志 3：实时监控、设备状态与参数配置页面完善",
        "module": "实时监控 / 设备状态 / 参数配置 / ECharts",
        "task_name": "完善监控、设备与参数配置页面，并保持与 dashboard 风格一致",
        "description": "在 dashboard 基础上继续完善实时监控、设备状态和参数配置页面。重点包括状态卡优化、设备健康展示、报警与运行数据呈现、参数表单结构和一致的浅色工控视觉语言。",
        "features": [
            "实时监控页优化 status-card，小卡片标题样式与仪表盘设备状态模块保持一致。",
            "设备状态页展示设备健康、状态、编号、告警和趋势类信息，图表统一使用 ECharts。",
            "参数配置页用于整理料盒、识别、动作和安全类参数，遵循 Element Plus 表单与校验语义。",
            "所有页面统一顶部状态栏、侧边导航、底部系统栏、浅色背景和紧凑信息布局。",
        ],
        "estimated": 7.0,
        "actual": 1.8,
        "traditional": [
            ["页面规格拆分", "1.2 h", "拆分实时监控、设备状态、参数配置三类页面的信息结构。"],
            ["组件与样式实现", "2.5 h", "实现状态卡、设备卡、表单区和表格区。"],
            ["图表与状态语义对齐", "1.5 h", "统一 ECharts、状态颜色、标题和字段样式。"],
            ["跨页面一致性修正", "1.8 h", "检查导航、顶部栏、模块标题和卡片间距是否一致。"],
        ],
        "prompts": [
            ["1", "用 ui-ux-pro-max 美化一下实时监控的 status-card。", "优化实时监控卡片结构、标题和状态展示。"],
            ["2", "所有小卡片的标题样式跟仪表盘的各设备状态模块卡片标题样式保持一致。", "抽齐跨页面小卡片标题风格，减少视觉割裂。"],
            ["3", "所有小卡片的标题上下间距收窄一些；再收窄。", "继续压缩卡片标题区域，提高监控页面的信息密度。"],
            ["4", "添加规则：图表使用 ECharts。", "将图表技术标准固化到 AGENTS.md 和页面实现约束中。"],
        ],
        "quality": "功能完整性 4.5 分；风格一致性 4.5 分；可扩展性 4.5 分。",
        "biggest_gain": "AI 对跨页面样式同步和状态语义统一提效明显，尤其适合把一个页面的卡片规格迁移到其他页面。",
        "prompt_pattern": "请把[页面B]中的[组件样式]与[页面A/样本组件]保持一致，并继续压缩[标题/间距/高度]，同时保留[业务字段]。",
        "suggestions": [
            "跨页面 UI 一致性最好指定样本模块，而不是只说“美化”。",
            "工控页面状态颜色要提前固化，避免不同页面对运行、报警、离线使用不同颜色。",
            "参数页要用 Element Plus 的表单、输入、校验和确认反馈来保证工业软件语义清晰。",
        ],
        "screenshots": [
            ("02-realtime.png", "图 2：实时监控页面截图"),
            ("04-parameters.png", "图 3：参数配置页面截图"),
            ("05-devices.png", "图 4：设备状态页面截图"),
        ],
        "code_assets": "src/widgets/realtime-monitor/*，src/widgets/device-health/*，src/widgets/parameter-config/*，src/shared/composables/useEChart.ts。",
        "screen_assets": "ai_log_render/screenshots/02-realtime.png，04-parameters.png，05-devices.png。",
        "verification": "已截图检查实时监控、参数配置、设备状态页面的主视图渲染。",
    },
    {
        "title": "日志 4：手动控制页面与分拣流程交互设计",
        "module": "手动控制 / 设备操作 / 分拣流程",
        "task_name": "新增手动控制页面，覆盖传送带、相机、夹爪和完整分拣流程",
        "description": "根据上位机人工调试场景新增手动控制页面，将传送带、相机、夹爪控制和完整分拣流程放在同一页，方便操作员在手动模式下进行单设备调试和完整流程验证。",
        "features": [
            "传送带控制包含图片、设备编号、当前状态、当前运行速度，以及启动、停止、复位故障按钮。",
            "相机控制包含图片、设备编号、当前状态、当前帧率，以及单次触发识别、查看识别结果按钮。",
            "夹爪控制包含图片、设备编号、当前状态、当前闭合程序、闭合程度滑块和百分比输入，并支持张开、闭合动作。",
            "分拣流程使用上料、传送、识别、抓取、放料、完成六步工作流，并展示本次分拣详细执行结果。",
            "工作流放在工作台图片上方，执行结果放在右侧，整体高度压缩以适配工作台布局。",
        ],
        "estimated": 7.5,
        "actual": 1.7,
        "traditional": [
            ["控制对象拆解", "1.2 h", "拆分传送带、相机、夹爪和流程控制的字段与动作。"],
            ["交互组件实现", "2.3 h", "实现按钮组、滑块、百分比输入、状态字段和禁用态。"],
            ["流程可视化设计", "2.0 h", "设计六步流程、工作台图片和结果面板布局。"],
            ["布局压缩与回归", "2.0 h", "根据页面高度反复压缩设备卡和流程区域。"],
        ],
        "prompts": [
            ["1", "添加一个手动控制页面：传送带控制、相机控制、夹爪控制、分拣流程。", "生成手动控制页的四大模块和设备操作按钮。"],
            ["2", "分拣流程：上料、传送、识别、抓取、放料、完成，工作流的表现方式。", "将分拣流程从普通描述改为六步工作流展示。"],
            ["3", "manual-device-card 内容高度可以再矮一些。", "压缩设备控制卡高度，提升一屏可见内容。"],
            ["4", "分拣流程放在图片的上面，执行结果可以放在本次分拣详细执行结果的右边。", "调整为上方流程、下方工作台、右侧结果的结构。"],
        ],
        "quality": "功能完整性 5 分；交互清晰度 4.5 分；可测试性 4.5 分。",
        "biggest_gain": "AI 能快速把自然语言设备控制需求转成页面结构和交互组件，减少手工设计控制面板的时间。",
        "prompt_pattern": "请新增[控制页面]，每个设备模块包含[图片/编号/状态/实时指标]，下方放[动作按钮]，并为[流程名称]设计[步骤列表]和[结果字段]。",
        "suggestions": [
            "设备控制类页面要先明确每个设备的实时字段和可执行动作。",
            "滑块和数字输入共用变量时，要在需求里直接说明，避免状态不同步。",
            "流程页面要把“步骤”和“结果”分开表达，操作员才能快速判断执行到哪一步。",
        ],
        "screenshots": [("03-manual.png", "图 5：手动控制页面截图")],
        "code_assets": "src/pages/manual/ManualPage.vue，src/widgets/manual-control/*。",
        "screen_assets": "ai_log_render/screenshots/03-manual.png。",
        "verification": "已截图检查手动控制页面进入手动模式后可见，六步流程和结果面板正常展示。",
    },
    {
        "title": "日志 5：国际化接入与 Vue 类型治理重构",
        "module": "Vue I18n / 类型组织 / 数据拆分 / 质量验证",
        "task_name": "接入中英文切换，并抽取 Vue SFC 中臃肿的 interface、静态数据和共享类型",
        "description": "在页面开发后继续补齐工程质量工作：默认中文，顶部使用图标加语言文本的单击切换；同时将各 Vue 文件中的复杂 interface、静态数据和展示模型抽取到 feature-local 的 types.ts、data.ts 或 shared/types 中。",
        "features": [
            "接入 Vue I18n，默认 zh-CN，支持中英文单击切换，交互为图标 + 中文，点击后变为图标 + English。",
            "仪表盘接入国际化，用户可见文案进入 locale 文件。",
            "将模块内非平凡 interface 抽取到相邻 types.ts，将静态展示数据、mock rows 和 i18n 源映射移入 data.ts/mock.ts。",
            "明确 entities 文件夹用于承载业务实体类型和纯业务模型，不放页面展示状态。",
            "补充 AGENTS.md 类型组织规则，防止后续 Vue SFC 再次堆积 interface 和大段数据。",
        ],
        "estimated": 6.5,
        "actual": 1.4,
        "traditional": [
            ["国际化资源整理", "1.5 h", "人工扫描用户可见文案并迁移到 zh-CN/en-US。"],
            ["语言切换交互实现", "1.0 h", "实现顶部语言切换按钮、Element Plus locale 和状态同步。"],
            ["类型与数据抽取", "2.5 h", "逐个 Vue 文件拆出 types.ts、data.ts 和 shared types。"],
            ["lint/build/test 验证", "1.5 h", "修正类型错误并确认构建、测试和 lint 通过。"],
        ],
        "prompts": [
            ["1", "加入国际化，默认中文；是不是要增加切换中英文的按钮？", "接入 Vue I18n，并在顶部状态栏加入语言切换入口。"],
            ["2", "中英文切换不要用下拉，就使用单击切换；图标 + 中文，点击后变成图标 + 英文。", "将下拉式语言选择改为单击切换的轻量交互。"],
            ["3", "仪表盘接入国际化。", "迁移 dashboard 用户可见文案到 locale 资源。"],
            ["4", "为什么我写了规则 Put shared types... 每个 vue 文件的 interface 都没有分模块写入 shared/types 里面？", "解释 shared/types 只放跨模块共享类型，feature 类型应优先放相邻 types.ts。"],
            ["5", "加入到规则，并优化；开始抽取吧。", "补充类型组织规则，并批量抽取各模块 interface 和静态数据。"],
        ],
        "quality": "功能完整性 4.5 分；类型可维护性 5 分；长期治理价值 5 分。",
        "biggest_gain": "AI 在跨文件类型抽取、规则解释和批量一致性检查上提效明显，能快速减少 Vue SFC 臃肿度。",
        "prompt_pattern": "请扫描[Vue 模块]，将超过[复杂度阈值]的 interface 抽到相邻 types.ts，将静态展示数据抽到 data.ts，并只把跨模块复用类型放进 shared/types。",
        "suggestions": [
            "不要把所有类型都塞进 shared/types，先判断类型是页面私有、feature 内复用，还是跨 feature 复用。",
            "国际化要优先迁移用户可见文案，路由标题、导航、按钮、状态和表格列名都要纳入范围。",
            "重构后要跑 lint、build 和 test，防止类型抽取造成导入路径或循环依赖问题。",
        ],
        "code_assets": "src/locales/*，src/widgets/*/types.ts，src/widgets/*/data.ts，src/shared/types/*，src/entities/*/types.ts。",
        "screen_assets": "语言切换入口位于顶部状态栏；可配合各页面截图归档。",
        "verification": "建议归档 pnpm run lint:fix、pnpm run build、pnpm run test 的通过记录。",
    },
]


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    title_run = title.add_run("《AI增效工程日志》")
    set_run_font(title_run, size=22, bold=True, color=COLORS["ink"])

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle_run = subtitle.add_run("ROBOFLOW 上位机机器人前端开发 · 汇总版")
    set_run_font(subtitle_run, size=11, color=COLORS["muted"])

    add_heading(doc, "功能整理总览", level=1)
    add_text(
        doc,
        "以下内容参考用户提供的《AI增效日志模板.docx》字段结构整理，覆盖本次对话中完成或规划落地的前端规范、页面功能、UI/UX 迭代、国际化和类型治理工作。",
        color=COLORS["muted"],
    )
    add_matrix(
        doc,
        ["功能/主题", "已整理内容", "可归档成果"],
        [
            ["项目规范与架构", "AGENTS.md、目录架构、路由规范、组件规范、前端-only 边界", "规则文档与开发规范"],
            ["仪表盘", "机器人状态、分拣指标、设备状态、快捷控制、历史任务、报警信息", "Dashboard 页面与截图"],
            ["实时/设备/参数", "实时监控状态卡、设备健康、参数配置、ECharts 与 Element Plus 交互", "三类页面与截图"],
            ["手动控制", "传送带、相机、夹爪、六步分拣流程和执行结果", "Manual 页面与截图"],
            ["工程治理", "Vue I18n、语言切换、types.ts/data.ts 抽取、shared/types 边界", "国际化与类型重构记录"],
        ],
        [3.6, 8.0, 5.0],
    )

    add_heading(doc, "截图总览", level=1)
    for image_name, caption in [
        ("01-dashboard.png", "图 A：仪表盘主界面"),
        ("02-realtime.png", "图 B：实时监控界面"),
        ("03-manual.png", "图 C：手动控制界面"),
    ]:
        add_screenshot(doc, image_name, caption)

    for index, log in enumerate(LOGS):
        if index > 0:
            doc.add_page_break()
        add_log(doc, log)

    doc.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    print(build())
